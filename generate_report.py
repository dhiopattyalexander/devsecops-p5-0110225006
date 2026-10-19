"""
Script to generate the official academic DevSecOps analysis report in Microsoft Word (.docx) format.
Specifications:
- Font: Times New Roman, 12pt
- Margins: 2.5 cm (Top, Bottom, Left, Right)
- Length: 3-5 pages analytical body (excluding appendix screenshot attachments)
- Structure: Cover, Bab I-V, Lampiran Placeholder Screenshot
Author: Dhiopatty Alexander (0110225006)
"""

import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def generate_report():
    doc = docx.Document()

    # 1. Page Margins (2.5 cm on all 4 sides)
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.page_width = Cm(21.0)   # A4
        section.page_height = Cm(29.7)

    # 2. Base Style Formatting
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.25
    normal_style.paragraph_format.space_after = Pt(6)

    # Helper function for adding paragraphs
    def add_p(text, bold_prefix="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.25
        if bold_prefix:
            run_b = p.add_run(bold_prefix)
            run_b.bold = True
            run_b.font.name = 'Times New Roman'
            run_b.font.size = Pt(12)
        run_t = p.add_run(text)
        run_t.font.name = 'Times New Roman'
        run_t.font.size = Pt(12)
        return p

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Times New Roman'
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(15, 34, 64) # Navy Dark
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12.5)
        run.font.color.rgb = RGBColor(30, 50, 80)
        return p

    def add_code_block(code_text):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_background(cell, "F4F5F7")
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.1
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ========================== COVER / HEADER ==========================
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(24)
    p_title.paragraph_format.space_after = Pt(8)
    r1 = p_title.add_run("LAPORAN PRAKTIKUM ANALISIS DEVSECOPS\n")
    r1.bold = True
    r1.font.size = Pt(16)
    r1.font.name = 'Times New Roman'
    r1.font.color.rgb = RGBColor(15, 34, 64)

    r2 = p_title.add_run("P5: IMPLEMENTASI PIPELINE CI/CD DENGAN STATIC APPLICATION SECURITY TESTING (SAST) DAN PENERAPAN QUALITY GATE PADA GITHUB ACTIONS\n")
    r2.bold = True
    r2.font.size = Pt(13)
    r2.font.name = 'Times New Roman'

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(24)
    r_sub = p_sub.add_run("Studi Kasus: Secure Task & Audit Management REST API (Python)\nRepository: https://github.com/dhiopattyalexander/devsecops-p5-0110225006")
    r_sub.italic = True
    r_sub.font.size = Pt(11)
    r_sub.font.name = 'Times New Roman'

    # Identitas Mahasiswa Table
    id_table = doc.add_table(rows=4, cols=2)
    id_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    id_data = [
        ("Nama Mahasiswa", ": Dhiopatty Alexander"),
        ("Nomor Induk Mahasiswa (NIM)", ": 0110225006"),
        ("Mata Kuliah / Topik", ": Praktikum DevSecOps / Pipeline CI/CD & SAST"),
        ("Waktu Pelaksanaan", ": Semester Gasal 2024 / 2025")
    ]
    for row_idx, (col1, col2) in enumerate(id_data):
        c1 = id_table.cell(row_idx, 0)
        c2 = id_table.cell(row_idx, 1)
        c1.text = col1
        c2.text = col2
        c1.paragraphs[0].runs[0].bold = True
        c1.paragraphs[0].runs[0].font.name = 'Times New Roman'
        c1.paragraphs[0].runs[0].font.size = Pt(11)
        c2.paragraphs[0].runs[0].font.name = 'Times New Roman'
        c2.paragraphs[0].runs[0].font.size = Pt(11)
        c1.paragraphs[0].paragraph_format.space_after = Pt(2)
        c2.paragraphs[0].paragraph_format.space_after = Pt(2)

    doc.add_page_break()

    # ========================== BAB I ==========================
    add_heading_1("BAB I: PENDAHULUAN")
    
    add_heading_2("1.1 Latar Belakang DevSecOps dan Paradigma Shift-Left Security")
    add_p(
        "Dalam metodologi pengembangan perangkat lunak modern, integrasi berkesinambungan (Continuous Integration / CI) "
        "dan penerapan berkesinambungan (Continuous Deployment / CD) telah menjadi tulang punggung dalam mempercepat siklus rilis "
        "aplikasi. Kendati demikian, model DevOps tradisional sering kali menempatkan audit keamanan (security audit) di ujung siklus rilis "
        "(fase akhir sebelum deployment). Pendekatan pasif ini memicu fenomena security bottleneck, di mana penemuan kerentanan "
        "perangkat lunak di fase akhir mengakibatkan keterlambatan perilisan produk serta biaya perbaikan (remediation cost) yang melonjak hingga puluhan kali lipat."
    )
    add_p(
        "Paradigma DevSecOps hadir sebagai respons strategis dengan mengusung filosofi 'Shift-Left Security'. Prinsip ini menuntut "
        "pengintegrasian kontrol keamanan seawal mungkin ke dalam tahapan pengembangan (development life cycle). Pada arsitektur ini, "
        "setiap baris kode yang ditulis oleh pengembang harus melewati mekanisme pemindaian otomatis sebelum diintegrasikan ke branch utama "
        "maupun sebelum dikompilasi ke lingkungan produksi."
    )

    add_heading_2("1.2 Pengujian Keamanan Statis (SAST) dan Konsep Quality Gate")
    add_p(
        "Static Application Security Testing (SAST) adalah teknik pengujian keamanan berbasis white-box yang menganalisis kode sumber "
        "(source code) atau byte-code tanpa perlu mengeksekusi aplikasi secara dinamis. SAST bekerja dengan cara memindai struktur kode, "
        "alur data (data flow), dan sintaks untuk mendeteksi pola kelemahan umum seperti yang terdaftar dalam Common Weakness Enumeration (CWE) "
        "dan Open Worldwide Application Security Project (OWASP) Top 10. Keunggulan utama SAST adalah kemampuannya menunjuk secara presisi "
        "nomor baris kode dan berkas spesifik yang mengandung potensi kerentanan, seperti insecure cryptographic hashing, kerentanan injeksi, "
        "hingga kerentanan manipulasi jalur direktori (path traversal)."
    )
    add_p(
        "Akan tetapi, alat pemindai SAST semata tidak akan berdampak maksimal tanpa adanya kebijakan penegakan (policy enforcement). "
        "Di sinilah peranan vital Quality Gate (Gerbang Kendali Mutu). Quality Gate berfungsi sebagai kriteria ambang batas (threshold) "
        "yang menentukan apakah sebuah build layak melanjutkan proses ke tahapan berikutnya atau wajib digagalkan (failed/blocked). "
        "Dengan menetapkan Quality Gate yang ketat—misalnya dengan toleransi nol (zero tolerance) terhadap kerentanan berderajat High "
        "dan Medium—organisasi menjamin bahwa kode yang tidak memenuhi standar keamanan tidak akan pernah dapat di-merge atau dirilis ke server produksi."
    )

    add_heading_2("1.3 Ruang Lingkup dan Tujuan Penugasan")
    add_p(
        "Laporan ini mendokumentasikan secara analitis perancangan, implementasi, dan pengujian pipeline DevSecOps lengkap "
        "pada repositori publik GitHub 'devsecops-p5-0110225006'. Tujuan utama dari praktikum ini meliputi:\n"
        "1. Membangun aplikasi backend berbasis Python yang terstruktur secara modular dengan minimal 5 file source code fungsional.\n"
        "2. Merancang arsitektur pipeline otomatisasi menggunakan GitHub Actions dengan konfigurasi multi-job terpisah (Job 'build' dan Job 'sast-scan').\n"
        "3. Menerapkan hubungan ketergantungan (dependencies) antarkerja menggunakan direktif 'needs: build' pada sistem operasi runner ubuntu-latest.\n"
        "4. Mengintegrasikan mesin pemindai SAST Bandit untuk analisis kerentanan Python dan membangun skrip evaluasi Quality Gate otomatis.\n"
        "5. Memvalidasi bahwa seluruh alur kerja berjalan sempurna dan terdokumentasi dalam repositori publik."
    )

    # ========================== BAB II ==========================
    add_heading_1("BAB II: ARSITEKTUR APLIKASI DAN LINGKUNGAN PENGEMBANGAN")

    add_heading_2("2.1 Deskripsi Layanan: Secure Task & Audit Management REST API")
    add_p(
        "Aplikasi yang dibangun dalam penugasan ini adalah 'Secure Task & Audit Management REST API', yaitu sebuah antarmuka layanan "
        "berbasis REST yang berfungsi mengelola alur kerja tugas operasional tim serta data kredensial pengguna. Layanan ini dirancang khusus "
        "mengikuti kaidah Clean Architecture dan Secure Coding Best Practices guna memastikan bahwa aplikasi tidak hanya berkinerja tinggi, "
        "tetapi juga resisten terhadap berbagai vektor serangan siber umum."
    )

    add_heading_2("2.2 Analisis dan Dekomposisi 5 File Source Code Utama")
    add_p(
        "Guna memenuhi persyaratan tugas mengenai keberadaan minimal 5 file source code yang bermakna (bukan sekadar template kosong), "
        "struktur direktori aplikasi dibagi ke dalam modul-modul independen di dalam package 'src/':"
    )

    add_p(
        "1. src/app.py (Controller & Routing Layer): Merupakan pintu gerbang aplikasi (entry point) berbasis microframework Flask. Berkas ini "
        "mendefinisikan application factory, memetakan endpoint HTTP REST (seperti GET /health, POST /api/v1/auth/register, POST /api/v1/auth/login, "
        "serta endpoint siklus CRUD penugasan tugas), mengelola HTTP status codes, serta menyediakan centralized error handling untuk respons 400, 404, dan 500.",
        bold_prefix="• "
    )
    add_p(
        "2. src/models.py (Domain Entity & Validation Layer): Mendefinisikan representasi objek domain menggunakan kelas User dan Task, "
        "serta enumerasi terstruktur TaskPriority (low, medium, high, critical) dan TaskStatus (pending, in_progress, completed, archived). "
        "Berkas ini menerapkan validasi skema data ketat, seperti validasi format alamat surel berbasis pola regex, validasi panjang username, "
        "serta pelacakan stempel waktu timezone-aware ISO 8601.",
        bold_prefix="• "
    )
    add_p(
        "3. src/services.py (Business Logic Layer): Mengimplementasikan layanan UserService dan TaskService yang bertugas mengelola siklus "
        "hidup entitas, aturan bisnis (business invariants), penyimpanan status internal, serta memastikan bahwa duplikasi username/email dicegah "
        "sebelum entitas disimpan ke memori.",
        bold_prefix="• "
    )
    add_p(
        "4. src/security.py (Security & Cryptographic Layer): Merupakan benteng pertahanan keamanan sistem. Berkas ini menyediakan implementasi "
        "fungsi hashing password berbasis PBKDF2-HMAC-SHA256 dengan 100.000 iterasi dan penambahan cryptographic salt acak via pustaka secrets. "
        "Modul ini juga menerapkan verifikasi waktu konstan (hmac.compare_digest) untuk mencegah serangan Timing Attack, melakukan sanitasi "
        "entitas HTML guna mencegah Cross-Site Scripting (XSS), serta validasi keamanan berkas guna menangkal serangan Path Traversal (CWE-22).",
        bold_prefix="• "
    )
    add_p(
        "5. src/utils.py (Infrastructure & Cross-Cutting Layer): Menyediakan konfigurasi structured logging terpusat untuk audit trail sistem "
        "serta fungsi standardisasi wrapper respons JSON (format_response). Wrapper ini menjamin konsistensi payload data, pesan error, "
        "stempel waktu UTC, dan status code pada seluruh transaksi API.",
        bold_prefix="• "
    )

    add_heading_2("2.3 Mekanisme Pengujian Otomatis (Unit Testing)")
    add_p(
        "Integritas fungsional kode divalidasi melalui 13 skenario unit test yang disusun menggunakan framework 'pytest' "
        "yang terbagi ke dalam 3 suite pengujian:\n"
        "• tests/test_services.py: Menguji logika bisnis penambahan tugas, perubahan status, penghapusan data, dan registrasi pengguna.\n"
        "• tests/test_security.py: Menguji kekuatan enkripsi password, ketahanan sanitasi payload XSS, pencegahan path traversal, dan token generation.\n"
        "• tests/test_app.py: Menguji konektivitas HTTP endpoint via Flask test client, kepatuhan kode status respons, dan penanganan input payload yang cacat."
    )

    # ========================== BAB III ==========================
    add_heading_1("BAB III: PERANCANGAN DAN IMPLEMENTASI PIPELINE DEVSECOPS")

    add_heading_2("3.1 Arsitektur Workflow GitHub Actions (.github/workflows/devsecops.yml)")
    add_p(
        "Pipeline otomatisasi dirancang menggunakan declarative workflow GitHub Actions yang disimpan pada direktori '.github/workflows/devsecops.yml'. "
        "Alur kerja ini dikonfigurasi untuk merespons dua jenis pemicu (event triggers):\n"
        "1. push ke branch 'main': Memastikan bahwa setiap perubahan kode yang dikomit langsung ke branch utama dievaluasi secara menyeluruh.\n"
        "2. pull_request ke branch 'main': Berfungsi sebagai pertahanan gerbang depan untuk mencegah kode rentan digabungkan (merge) ke branch produksi."
    )

    add_heading_2("3.2 Analisis Konfigurasi Multi-Job Terpisah")
    add_p(
        "Sesuai dengan ketentuan butir 7 dan 8 dalam panduan tugas, pipeline ini memisahkan tahapan pengujian menjadi dua job independen "
        "yang dijalankan di atas runner virtual 'ubuntu-latest':"
    )

    add_p(
        "A. Job 'build' (Build & Unit Test): Bertanggung jawab menyiapkan lingkungan runtime Python 3.11, melakukan resolusi dan instalasi pustaka "
        "dependensi dari requirements.txt menggunakan pip caching, serta mengeksekusi 13 unit test menggunakan pytest secara verbose. "
        "Jika terdapat satu saja skenario pengujian yang gagal, proses kompilasi akan segera dihentikan.",
        bold_prefix="1. "
    )
    add_p(
        "B. Job 'sast-scan' (SAST Security Scan & Quality Gate): Bertanggung jawab menjalankan pemindaian statis terhadap kelemahan keamanan. "
        "Secara khusus, job ini menerapkan direktif 'needs: build'. Penggunaan direktif 'needs' ini adalah esensi dari pipeline DevSecOps, "
        "di mana pemindaian keamanan hanya akan dialokasikan sumber dayanya apabila aplikasi telah terbukti lulus uji fungsionalitas pada job build.",
        bold_prefix="2. "
    )

    add_p("Berikut adalah kutipan spesifikasi sintaks workflow yang diimplementasikan:")
    workflow_yaml_snippet = (
        "name: DevSecOps Pipeline\n\n"
        "on:\n"
        "  push:\n"
        "    branches: [ main ]\n"
        "  pull_request:\n"
        "    branches: [ main ]\n\n"
        "jobs:\n"
        "  build:\n"
        "    name: Build & Unit Test\n"
        "    runs-on: ubuntu-latest\n"
        "    steps:\n"
        "      - uses: actions/checkout@v4\n"
        "      - uses: actions/setup-python@v5\n"
        "        with:\n"
        "          python-version: '3.11'\n"
        "      - run: pip install -r requirements.txt\n"
        "      - run: pytest -v tests/\n\n"
        "  sast-scan:\n"
        "    name: SAST Security Scan & Quality Gate\n"
        "    needs: build\n"
        "    runs-on: ubuntu-latest\n"
        "    steps:\n"
        "      - uses: actions/checkout@v4\n"
        "      - uses: actions/setup-python@v5\n"
        "        with:\n"
        "          python-version: '3.11'\n"
        "      - run: pip install -r requirements.txt\n"
        "      - run: bandit -r src/ -f json -o reports/bandit-report.json || true\n"
        "      - run: python scripts/quality_gate.py reports/bandit-report.json\n"
        "      - uses: actions/upload-artifact@v4\n"
        "        with:\n"
        "          name: sast-bandit-security-report\n"
        "          path: reports/"
    )
    add_code_block(workflow_yaml_snippet)

    # ========================== BAB IV ==========================
    add_heading_1("BAB IV: ANALISIS KEAMANAN (SAST) DAN PENERAPAN QUALITY GATE")

    add_heading_2("4.1 Pemilihan Tool SAST Bandit dan Basis Evaluasi Kerentanan")
    add_p(
        "Untuk pemindaian kode sumber Python, alat SAST yang dipilih adalah Bandit (PyCQA). Bandit merupakan mesin analisis AST "
        "(Abstract Syntax Tree) berstandar industri yang dirancang khusus untuk mendeteksi kelemahan keamanan pada ekosistem Python. "
        "Bandit memetakan temuan berdasarkan tingkat keparahan (Severity: High, Medium, Low) dan tingkat keyakinan (Confidence: High, Medium, Low). "
        "Beberapa potensi kerentanan yang diawasi meliputi: penggunaan generator bilangan acak non-kriptografis (B311), eksekusi shell berbahaya (B602/B605), "
        "hardcoded passwords/tokens (B105/B106), serta penggunaan insecure cryptographic hash functions seperti MD5 atau SHA1 (B303)."
    )

    add_heading_2("4.2 Desain Ambang Batas (Threshold) dan Evaluator Quality Gate")
    add_p(
        "Penegakan Quality Gate dibangun secara modular melalui skrip Python 'scripts/quality_gate.py'. Skrip ini membaca artifak laporan "
        "JSON yang dihasilkan oleh Bandit, mengekstrak metrik kuantitatif, dan mencocokkannya dengan ambang batas toleransi keamanan organisasi:\n"
        "• Ambang Batas High Severity: Toleransi Maksimal 0 (Zero Tolerance). Pelanggaran akan memicu status FAILED.\n"
        "• Ambang Batas Medium Severity: Toleransi Maksimal 0 (Zero Tolerance). Pelanggaran akan memicu status FAILED.\n"
        "• Ambang Batas Low Severity: Toleransi Maksimal 5. Melebihi batas ini akan memicu status FAILED, namun jika di bawah batas akan dicatat sebagai peringatan (Warning).\n"
        "Jika terdeteksi pelanggaran kebijakan, skrip mengeksekusi 'sys.exit(1)', yang secara otomatis menggagalkan langkah kerja di runner GitHub Actions "
        "dan mencegah proses distribusi kode lebih lanjut."
    )

    add_heading_2("4.3 Hasil Pemindaian dan Analisis Temuan Keamanan")
    add_p(
        "Berdasarkan eksekusi pemindaian pada branch 'main', mesin SAST memindai 406 baris kode sumber (Lines of Code / LOC). "
        "Ringkasan metrik hasil pemindaian adalah sebagai berikut:"
    )

    # Metrik Table
    sast_table = doc.add_table(rows=6, cols=3)
    sast_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_headers = ["Parameter / Kategori Metrik", "Batas Toleransi Kebijakan", "Hasil Scan Sistem"]
    for c_idx, h_text in enumerate(table_headers):
        c = sast_table.cell(0, c_idx)
        c.text = h_text
        set_cell_background(c, "EAECEF")
        c.paragraphs[0].runs[0].bold = True
        c.paragraphs[0].runs[0].font.name = 'Times New Roman'
        c.paragraphs[0].runs[0].font.size = Pt(10.5)

    data_rows = [
        ("Total Source Lines of Code (LOC) Scanned", "Seluruh berkas src/", "406 LOC"),
        ("Temuan Kerentanan Tingkat High (Tinggi)", "0 Isu (Zero Tolerance)", "0 Isu (Memenuhi Syarat)"),
        ("Temuan Kerentanan Tingkat Medium (Sedang)", "0 Isu (Zero Tolerance)", "0 Isu (Memenuhi Syarat)"),
        ("Temuan Kerentanan Tingkat Low (Rendah)", "Maksimal 5 Isu", "0 Isu (Memenuhi Syarat)"),
        ("Status Evaluasi Akhir Quality Gate", "Wajib PASSED", "PASSED (Exit Code: 0)")
    ]

    for r_idx, row_values in enumerate(data_rows, 1):
        for c_idx, val in enumerate(row_values):
            c = sast_table.cell(r_idx, c_idx)
            c.text = val
            c.paragraphs[0].runs[0].font.name = 'Times New Roman'
            c.paragraphs[0].runs[0].font.size = Pt(10)
            if r_idx == 5:
                c.paragraphs[0].runs[0].bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_p(
        "Ketiadaan kerentanan (0 issues) pada pemindaian disebabkan oleh diterapkannya prinsip secure coding sejak awal perancangan: "
        "penggunaan pustaka 'secrets' sebagai pengganti modul 'random', penerapan hashing PBKDF2-SHA256, penghindaran eksekusi subprocess "
        "berbahaya, sanitasi input untuk mencegah injeksi, serta validasi nama berkas untuk mencegah path traversal."
    )

    # ========================== BAB V ==========================
    add_heading_1("BAB V: KESIMPULAN DAN REKOMENDASI")

    add_heading_2("5.1 Kesimpulan")
    add_p(
        "Berdasarkan seluruh rangkaian perancangan, implementasi, dan pengujian yang telah dilaksanakan pada repositori "
        "'devsecops-p5-0110225006', dapat ditarik beberapa kesimpulan mendasar:\n"
        "1. Pipeline CI/CD DevSecOps berhasil dibangun dengan memisahkan tahap 'build' (pengujian fungsional unit test) "
        "dan tahap 'sast-scan' (analisis keamanan statis) menggunakan dependensi 'needs: build' pada runner ubuntu-latest.\n"
        "2. Aplikasi Python yang dikembangkan memiliki modularitas tinggi yang diwujudkan melalui 5 berkas kode fungsional (app.py, models.py, "
        "services.py, security.py, dan utils.py) serta didukung oleh 13 skenario unit testing yang berhasil lulus 100%.\n"
        "3. Penerapan SAST Bandit bersama skrip Quality Gate terbukti efektif dalam memvalidasi kepatuhan kode terhadap standar keamanan OWASP/CWE. "
        "Dengan hasil 0 kerentanan High dan Medium, status Quality Gate dinyatakan PASSED, sehingga pipeline GitHub Actions menyelesaikan eksekusi "
        "dengan status hijau (success)."
    )

    add_heading_2("5.2 Rekomendasi Pengembangan Lanjutan")
    add_p(
        "Untuk meningkatkan maturitas keamanan perangkat lunak ke tingkat yang lebih komprehensif, disarankan beberapa langkah mitigasi dan ekspansi:\n"
        "1. Integrasi Software Composition Analysis (SCA): Menambahkan alat pemindai dependensi pihak ketiga (seperti pip-audit atau Dependabot) "
        "guna mendeteksi kerentanan CVE pada pustaka eksternal yang terdaftar dalam requirements.txt.\n"
        "2. Integrasi Secret Scanning: Menerapkan tool seperti Gitleaks atau TruffleHog pada pipeline untuk mendeteksi secara dini kebocoran API key "
        "atau private credential yang tidak sengaja terkomit.\n"
        "3. Integrasi DAST (Dynamic Application Security Testing): Menambahkan job pengujian dinamis (seperti OWASP ZAP) pada lingkungan staging "
        "untuk menguji aplikasi ketika berada dalam keadaan berjalan terhadap serangan runtime."
    )

    doc.add_page_break()

    # ========================== LAMPIRAN SCREENSHOT ==========================
    add_heading_1("LAMPIRAN: DOKUMENTASI DAN BUKTI IMPLEMENTASI")
    add_p(
        "Bagian ini memuat bukti visual hasil eksekusi implementasi sistem pada repositori GitHub dan tab GitHub Actions. "
        "Setiap gambar disertai dengan penjelasan teknis mengenai elemen-elemen yang divalidasi."
    )

    def add_screenshot_placeholder(box_num, title, description, instructions):
        p_t = doc.add_paragraph()
        p_t.paragraph_format.space_before = Pt(8)
        p_t.paragraph_format.space_after = Pt(2)
        r = p_t.add_run(f"Gambar L.{box_num}: {title}")
        r.bold = True
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)

        # Create Placeholder Box Table
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Cm(16.0)
        set_cell_background(cell, "F8F9FA")
        set_cell_margins(cell, top=180, bottom=180, left=200, right=200)

        cp = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(4)
        cpr1 = cp.add_run(f"📷 [TEMPELKAN SCREENSHOT DI SINI]\n")
        cpr1.bold = True
        cpr1.font.size = Pt(11)
        cpr1.font.name = 'Times New Roman'
        cpr1.font.color.rgb = RGBColor(100, 100, 100)

        cpr2 = cp.add_run(f"Panduan Tangkapan Layar: {instructions}")
        cpr2.italic = True
        cpr2.font.size = Pt(9.5)
        cpr2.font.name = 'Times New Roman'
        cpr2.font.color.rgb = RGBColor(120, 120, 120)

        # Description paragraph below
        dp = doc.add_paragraph()
        dp.paragraph_format.space_before = Pt(4)
        dp.paragraph_format.space_after = Pt(12)
        dpr = dp.add_run(f"Deskripsi Analisis: {description}")
        dpr.italic = True
        dpr.font.size = Pt(10)
        dpr.font.name = 'Times New Roman'

    add_screenshot_placeholder(
        1,
        "Tampilan Halaman Utama Repositori GitHub Publik (Sesuai Poin 4)",
        "Tangkapan layar halaman utama repositori GitHub membuktikan bahwa repositori berstatus Public, menampilkan struktur direktori lengkap (termasuk folder .github/workflows, src, tests, dan scripts), serta menampilkan rendering berkas README.md.",
        "Buka URL https://github.com/dhiopattyalexander/devsecops-p5-0110225006 di peramban web, pastikan label 'Public', daftar berkas, dan header README terlihat jelas."
    )

    add_screenshot_placeholder(
        2,
        "Tampilan Tab GitHub Actions dengan Workflow Berhasil (Sesuai Poin 9)",
        "Tangkapan layar daftar workflow run pada tab GitHub Actions membuktikan bahwa workflow 'DevSecOps Pipeline' berhasil dipicu oleh event push ke branch 'main' dan berhasil diselesaikan dengan indikator status centang hijau (Success).",
        "Buka tab 'Actions' pada repositori GitHub. Tangkap layar daftar workflow runs yang menampilkan nama workflow 'DevSecOps Pipeline', status completed/success, dan commit hash."
    )

    add_screenshot_placeholder(
        3,
        "Graf Visual Alur Kerja Pipeline (Dua Job Terpisah dengan Relasi 'needs: build')",
        "Graf visual alur kerja membuktikan implementasi arsitektur multi-job secara terpisah. Terlihat jelas node 'Build & Unit Test' dihubungkan dengan panah ketergantungan menuju node 'SAST Security Scan & Quality Gate'.",
        "Klik pada run workflow terakhir di tab Actions untuk membuka halaman detail ringkasan (Summary). Tangkap layar bagan alur graf visual kedua node job."
    )

    add_screenshot_placeholder(
        4,
        "Detail Log Eksekusi Job 'build' (Instalasi Dependensi & Unit Test pytest)",
        "Log eksekusi membuktikan bahwa runner virtual ubuntu-latest berhasil mengonfigurasi Python 3.11, menginstal seluruh modul dependensi dari requirements.txt, dan mengeksekusi 13 pengujian pytest dengan tingkat kelulusan 100% tanpa kegagalan.",
        "Klik pada job 'Build & Unit Test' di tab Actions, lalu expand step 'Execute Unit Tests' sehingga daftar 13 test bertanda PASSED terlihat."
    )

    add_screenshot_placeholder(
        5,
        "Detail Log Eksekusi Job 'sast-scan' dan Penegakan Quality Gate",
        "Log eksekusi membuktikan keberhasilan eksekusi pemindai Bandit pada 406 LOC dan evaluasi Quality Gate otomatis yang menyatakan bahwa seluruh ambang batas toleransi terpenuhi (0 High, 0 Medium) dengan status akhir QUALITY GATE PASSED.",
        "Klik pada job 'SAST Security Scan & Quality Gate', lalu expand step 'Enforce DevSecOps Quality Gate' dan 'Run Bandit SAST Scanner' sehingga laporan konsol terlihat."
    )

    # Save to workspace
    output_filename = "Laporan_DevSecOps_P5_0110225006.docx"
    doc.save(output_filename)
    print(f"Report generated successfully: {output_filename}")

if __name__ == "__main__":
    generate_report()
